import sys
import subprocess

# Auto-install python-docx if missing
try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx'])
    import docx

from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()

# 1. Setup Document Styling (Times New Roman, 12pt, 1.5 spacing)
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.line_spacing = 1.5

with open('CACS452_Project_Report.md', 'r', encoding='utf-8') as f:
    lines = f.readlines()

# 2. Parse Markdown and write to Word
for line in lines:
    line = line.strip()
    if not line:
        continue
    
    if line.startswith('### '):
        heading = doc.add_heading(line[4:].replace('**', ''), level=3)
    elif line.startswith('## '):
        heading = doc.add_heading(line[3:].replace('**', ''), level=2)
    elif line.startswith('# '):
        heading = doc.add_heading(line[2:].replace('**', ''), level=1)
        heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    elif line.startswith('---'):
        doc.add_page_break()
    elif line.startswith('* '):
        doc.add_paragraph(line[2:].replace('**', ''), style='List Bullet')
    else:
        # Standard paragraph (strip rough markdown)
        clean_text = line.replace('**', '').replace('*', '')
        doc.add_paragraph(clean_text)

# 3. Save
doc.save('CACS452_Project_Report.docx')
print("Successfully generated CACS452_Project_Report.docx!")
